"""Editable visual rendering and unchanged-source-slide regression checks."""
import io
import zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
import pptx_structure_edit as edit


def main():
    plans = [
        {"kind": "process", "columns": ["工程", "内容"], "rows": [["回収", "回収する"], ["処理", "選別する"], ["処分", "処分する"]]},
        {"kind": "classification", "columns": ["分類", "内容"], "rows": [["分類A", "内容A"], ["分類B", "内容B"]]},
        {"kind": "comparison", "columns": ["比較軸", "地域企業", "大手"], "rows": [["顧客", "地域", "広域"], ["設備", "回収車", "処理設備"]]},
        {"kind": "bar", "columns": ["企業", "2025年度売上高（億円）"], "rows": [["試験A", "120"], ["試験B", "80"]], "values": [120, 80]},
    ]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    original = prs.slides.add_slide(prs.slide_layouts[6])
    title = original.shapes.add_textbox(Inches(.5), Inches(.3), Inches(12), Inches(.6))
    title.text = "Original"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    body = original.shapes.add_textbox(Inches(.5), Inches(2), Inches(10), Inches(2))
    body.text = "Source content remains unchanged"
    before = io.BytesIO()
    prs.save(before)
    for plan in plans:
        warnings = []
        index = edit.op_add_slide(prs, {"title": plan["kind"], "kind": "table", "table": {"columns": plan["columns"], "rows": plan["rows"]}, "composition": plan}, warnings)
        slide = prs.slides[index]
        assert any(sh.name.startswith("AzureChat Visual ") for sh in slide.shapes)
        text = "\n".join(sh.text for sh in slide.shapes if sh.has_text_frame)
        if plan["kind"] == "bar":
            chart = next(sh.chart for sh in slide.shapes if sh.has_chart)
            assert list(chart.series[0].values) == plan["values"]
            assert [c.label for c in chart.plots[0].categories] == [r[0] for r in plan["rows"]]
        else:
            assert all(cell in text for row in plan["rows"] for cell in row)
        assert all(sh.left >= 0 and sh.top >= 0 and sh.left + sh.width <= prs.slide_width + 2 and sh.top + sh.height <= prs.slide_height + 2 for sh in slide.shapes)
    for op in [
        {"title":"企業紹介","kind":"table","keyMessage":"主要企業の紹介","table":{"columns":["企業","2025年度売上高（億円）"],"rows":[["試験A","120"],["試験B","80"]]}},
        {"title":"主要指標","kind":"stats","stats":[{"label":"売上","value":"120","period":"2025年度","unit":"億円"},{"label":"件数","value":"80","period":"2025年度","unit":"件"}]},
        {"title":"企業比較","kind":"two_column","bullets":["回収サービス","処理サービス","地域で事業展開"],"table":{"columns":["企業","特徴"],"rows":[["試験A","回収"],["試験B","処理"]]}},
    ]:
        warnings=[]
        index=edit.op_add_slide(prs,op,warnings)
        slide=prs.slides[index]
        assert not any(sh.shape_type == 1 and sh.height == Pt(4) for sh in slide.shapes), "No decorative stripe"
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.name.startswith("AzureChat"): continue
            basis=shape.name in ("AzureChat Metric Basis","AzureChat Stat Basis","AzureChat Qualifications","AzureChat Added Sources")
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    assert run.font.size.pt >= (10 if basis else 14), (shape.name,run.text,run.font.size.pt)
        if op["title"] == "企業紹介":
            assert any(sh.name == "AzureChat Company Metric" for sh in slide.shapes)
            basis_text = [sh.text for sh in slide.shapes if sh.has_text_frame and sh.name == "AzureChat Metric Basis"]
            assert basis_text == ["2025年度売上高（億円）", "2025年度売上高（億円）"]
    after = io.BytesIO()
    prs.save(after)
    with zipfile.ZipFile(before) as a, zipfile.ZipFile(after) as b:
        assert a.read("ppt/slides/slide1.xml") == b.read("ppt/slides/slide1.xml")
    import sys
    if len(sys.argv) > 1:
        prs.save(sys.argv[1])
    print("PASS Python visual compositions: all source cells, editable shapes, bounds and original slide XML")


if __name__ == "__main__":
    main()
